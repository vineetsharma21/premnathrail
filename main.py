import math
import io 
import os
import subprocess
from fastapi import FastAPI, HTTPException, Depends, status, Request
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from typing import Dict, Any, List, Optional
from jinja2 import Environment, FileSystemLoader

# --- Database & Security Imports ---
from sqlalchemy.orm import Session
import bcrypt  # Direct import used now
import models   
import schemas  
from database import engine, get_db 

# --- Library Imports ---
try:
    import docx
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    docx = None

# --- Import Calculation Logic ---
from Hydraulic_Motor_Calculation import Calculator as HydraulicCalculator
from QmaxCalculator_Logic import QmaxCalculatorLogic, SIGMA_B_OPTIONS
from LoadDistribution_Logic import perform_calculations as perform_load_distro_calc
from LoadDistribution_Logic import format_detailed_steps as format_load_distro_steps
from LoadDistribution_Logic import create_report_docx as create_load_distro_docx
from Tractive_Effort_Logic import perform_te_calculations, format_te_report_text, create_te_report_docx
from Vehicle_Performance_Logic import VehiclePerformanceCalculator

# NEW: Braking Logic Import
from BrakingWebLogic import BrakingCalculator

# --- Setup ---
from contextlib import asynccontextmanager

@asynccontextmanager
async def lifespan(app: FastAPI):
    # Startup
    print("FastAPI application starting up...")
    print(f"Current working directory: {os.getcwd()}")
    print(f"Files in directory: {os.listdir('.')}")
    yield
    # Shutdown (if needed)

app = FastAPI(
    title="Engineering Calculator API",
    description="Advanced calculations for Hydraulic, Qmax, Load Distribution, Tractive Effort, Vehicle Performance and Braking Analysis",
    version="1.0",
    lifespan=lifespan
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"], 
    allow_credentials=True, 
    allow_methods=["*"], 
    allow_headers=["*"], 
)

# Initialize database tables
try:
    models.Base.metadata.create_all(bind=engine)
    print("Database tables created successfully")
except Exception as e:
    print(f"Database initialization warning: {e}")

# Mount static files
app.mount("/static", StaticFiles(directory="."), name="static")

@app.get("/health")
async def health_check():
    return {"status": "healthy", "message": "Engineering Calculator API is running"}

# --- NEW SECURITY SETUP (Using bcrypt directly) ---
def get_password_hash(password):
    """Hashes a password using bcrypt."""
    pwd_bytes = password.encode('utf-8')
    salt = bcrypt.gensalt()
    hashed = bcrypt.hashpw(pwd_bytes, salt)
    return hashed.decode('utf-8')

def verify_password(plain_password, hashed_password):
    """Verifies a password against its hash."""
    try:
        pwd_bytes = plain_password.encode('utf-8')
        hash_bytes = hashed_password.encode('utf-8')
        return bcrypt.checkpw(pwd_bytes, hash_bytes)
    except Exception:
        return False

# --- Calculators ---
hydraulic_calculator = HydraulicCalculator()
qmax_calculator = QmaxCalculatorLogic()
braking_calculator = BrakingCalculator()

# --- Pydantic Models ---
class HydraulicRawInput(BaseModel):
    calc_mode: str
    weight: str
    axles: str
    speed: str
    max_vehicle_rpm: str
    pto_gear_ratio: str
    engine_gear_ratio: str
    axle_gear_box_ratio: str
    slope_percent: str
    curve_degree: str
    wheel_diameter: str
    num_motors: str
    per_axle_motor: str
    pressure: str
    mech_eff_motor: str
    vol_eff_motor: str
    motor_disp_in: str
    max_motor_rpm: str
    vol_eff_pump: str
    pump_disp_in: str
    max_pump_rpm: str

class QmaxRawInput(BaseModel):
    d: str
    sigma_b_selection: str
    sigma_b_custom: str
    v_head: str

class LoadDistroRawInput(BaseModel):
    config_type: str
    total_load: str
    front_percent: str
    q1_percent: str
    q3_percent: str

class TractiveEffortRawInput(BaseModel):
    load: str
    loco_weight: str
    gradient: str
    curvature: str
    speed: str
    mode: str
    grad_type: str
    curvature_unit: str

class VehiclePerformanceRawInput(BaseModel):
    max_curve: str
    max_slope: str
    loco_gvw: str
    max_speed: str
    num_axles: str
    rear_axle_ratio: str
    gear_ratios: str
    shunting_load: str
    peak_power: str
    friction_mu: str
    wheel_dia: str
    min_rpm: str
    max_rpm: str
    torque_curve: Dict[str, float] = Field(default_factory=dict) 

class BrakingWebInput(BaseModel):
    doc_no: str = ""
    made_by: str = ""
    checked_by: str = ""
    approved_by: str = ""
    mass_kg: float
    speed_kmh: float
    number_of_wheels: int
    wheel_dia: float
    speed_increment: float = 10.0
    gradient_increment: int = 5
    
    # Rail parameters
    speed_input: str = ""  # Comma-separated speeds
    gradient_input_str: str = ""  # Comma-separated gradients
    gradient_input: float = 0.0  # Max gradient for calculations
    gradient_type: str = "1 in G"
    
    # Scenario selection
    scenario_straight: bool = True
    scenario_moving_up: bool = True
    scenario_moving_down: bool = True
    
    # Road parameters
    road_mode_enabled: bool = False
    road_speeds: str = ""
    road_gradients: str = ""
    road_gradient_type: str = "Percentage (%)"
    friction: float = 0.7
    
    # Distance source
    distance_source: str = "EN Standard"
    custom_distance: float = 0.0
    
    # GBR option
    include_gbr: bool = True


# --- Helper Validation ---
def _validate_input(value_str, type_func, name):
    if not value_str: raise ValueError(f"'{name}' cannot be empty.")
    try: return type_func(value_str)
    except: raise ValueError(f"Invalid value for '{name}'")

def process_and_validate_hydraulic_inputs(raw: HydraulicRawInput):
    inputs = raw.dict()
    try:
        inputs['weight'] = float(raw.weight)
        inputs['axles'] = int(float(raw.axles))
        inputs['speed'] = float(raw.speed)
        inputs['max_vehicle_rpm'] = float(raw.max_vehicle_rpm)
        inputs['pto_gear_ratio'] = float(raw.pto_gear_ratio)
        inputs['engine_gear_ratio_list'] = [float(x) for x in (raw.engine_gear_ratio or '').split(',') if x.strip()]
        inputs['axle_gear_box_ratio'] = float(raw.axle_gear_box_ratio)
        inputs['slope_percent'] = float(raw.slope_percent)
        inputs['curve_degree'] = float(raw.curve_degree)
        inputs['wheel_diameter'] = float(raw.wheel_diameter)
        inputs['num_motors'] = int(float(raw.num_motors))
        inputs['per_axle_motor'] = int(float(raw.per_axle_motor))
        inputs['pressure'] = float(raw.pressure)
        inputs['mech_eff_motor'] = float(raw.mech_eff_motor)
        inputs['vol_eff_motor'] = float(raw.vol_eff_motor)
        inputs['motor_disp_in'] = float(raw.motor_disp_in)
        inputs['max_motor_rpm'] = float(raw.max_motor_rpm)
        inputs['vol_eff_pump'] = float(raw.vol_eff_pump)
        inputs['pump_disp_in'] = float(raw.pump_disp_in)
        inputs['max_pump_rpm'] = float(raw.max_pump_rpm)
    except Exception as e:
        raise ValueError(f"Invalid hydraulic input: {e}")
    return inputs, raw.dict()

def process_and_validate_qmax_inputs(raw: QmaxRawInput):
    inputs = raw.dict()
    if raw.sigma_b_selection == "Custom": inputs['sigma_b'] = float(raw.sigma_b_custom)
    else: inputs['sigma_b'] = SIGMA_B_OPTIONS.get(raw.sigma_b_selection, 0)
    inputs['d'] = float(raw.d)
    inputs['v_head'] = float(raw.v_head)
    return inputs, raw.dict()

def process_and_validate_load_distro_inputs(raw: LoadDistroRawInput):
    inputs = raw.dict()
    inputs['total_load'] = float(raw.total_load)
    inputs['front_percent'] = float(raw.front_percent)
    inputs['q1_percent'] = float(raw.q1_percent)
    inputs['q3_percent'] = float(raw.q3_percent)
    return inputs, raw.dict()

def process_and_validate_te_inputs(raw: TractiveEffortRawInput):
    i = raw.dict()
    i['load'] = float(raw.load); i['loco_weight'] = float(raw.loco_weight)
    i['gradient'] = float(raw.gradient); i['curvature'] = float(raw.curvature)
    i['speed'] = float(raw.speed)
    return i, raw.dict()

def process_and_validate_vehicle_performance_inputs(raw: VehiclePerformanceRawInput):
    inputs = raw.dict()
    try:
        inputs['max_curve'] = float(raw.max_curve) if raw.max_curve not in (None, "") else 0.0
    except Exception:
        inputs['max_curve'] = 0.0
    try:
        inputs['max_slope'] = float(raw.max_slope)
    except Exception:
        inputs['max_slope'] = 0.0
    try:
        inputs['loco_gvw_kg'] = float(raw.loco_gvw)
    except Exception:
        inputs['loco_gvw_kg'] = 0.0
    try:
        inputs['max_speed_kmh'] = float(raw.max_speed)
    except Exception:
        inputs['max_speed_kmh'] = 0.0
    try:
        inputs['num_axles'] = int(float(raw.num_axles))
    except Exception:
        inputs['num_axles'] = 1
    try:
        inputs['rear_axle_ratio'] = float(raw.rear_axle_ratio)
    except Exception:
        inputs['rear_axle_ratio'] = 1.0
    try:
        inputs['gear_ratios'] = [float(x) for x in (raw.gear_ratios or '').split(',') if x.strip()]
        if not inputs['gear_ratios']: inputs['gear_ratios'] = [1.0]
    except Exception:
        inputs['gear_ratios'] = [1.0]
    try:
        inputs['shunting_load_t'] = float(raw.shunting_load)
    except Exception:
        inputs['shunting_load_t'] = 0.0
    try:
        inputs['peak_power_kw'] = float(raw.peak_power)
    except Exception:
        inputs['peak_power_kw'] = 0.0
    try:
        inputs['friction_mu'] = float(raw.friction_mu)
    except Exception:
        inputs['friction_mu'] = 0.3
    try:
        inputs['wheel_dia_m'] = float(raw.wheel_dia)
    except Exception:
        inputs['wheel_dia_m'] = 0.73
    try:
        inputs['min_rpm'] = int(float(raw.min_rpm))
    except Exception:
        inputs['min_rpm'] = 100
    try:
        inputs['max_rpm'] = int(float(raw.max_rpm))
    except Exception:
        inputs['max_rpm'] = 2500
    # Normalize torque_curve keys/values to int->float
    torque_curve = {}
    try:
        for k, v in raw.torque_curve.items():
            try:
                kk = int(float(k))
            except Exception:
                kk = int(k) if isinstance(k, int) else None
            if kk is None: continue
            torque_curve[kk] = float(v)
    except Exception:
        torque_curve = {}
    inputs['torque_curve'] = torque_curve
    return inputs, raw.dict()


# ==========================================
#      AUTH & LICENSE ENDPOINTS
# ==========================================

@app.post("/signup")
def signup(user: schemas.UserCreate, db: Session = Depends(get_db)):
    db_user = db.query(models.User).filter(models.User.email == user.email).first()
    if db_user: raise HTTPException(status_code=400, detail="Email already registered")
    hashed_password = get_password_hash(user.password)
    new_user = models.User(email=user.email, hashed_password=hashed_password)
    db.add(new_user); db.commit(); db.refresh(new_user)
    return {"message": "Account created successfully", "user_id": new_user.id}

@app.post("/login")
def login(user: schemas.UserLogin, db: Session = Depends(get_db)):
    db_user = db.query(models.User).filter(models.User.email == user.email).first()
    if not db_user or not verify_password(user.password, db_user.hashed_password):
        raise HTTPException(status_code=400, detail="Invalid email or password")
    return {"message": "Login successful", "user_id": db_user.id, "email": db_user.email, "is_license_active": db_user.is_license_active}

@app.post("/activate_license")
def activate_license(activation: schemas.LicenseActivate, user_id: int, db: Session = Depends(get_db)):
    user = db.query(models.User).filter(models.User.id == user_id).first()
    if not user: raise HTTPException(status_code=404, detail="User not found")
    license_entry = db.query(models.LicenseKey).filter(models.LicenseKey.key == activation.license_key).first()
    if not license_entry: raise HTTPException(status_code=400, detail="Invalid License Key")
    if license_entry.is_used: raise HTTPException(status_code=400, detail="This License Key is already used")
    license_entry.is_used = True; user.is_license_active = True; db.commit()
    return {"message": "License Activated!"}

# ==========================================
#      STATIC PAGES (HTML SERVING)
# ==========================================

@app.get("/")
async def serve_home(): return FileResponse('index.html')

@app.get("/login")
async def serve_login():
    if os.path.exists("login.html"): return FileResponse('login.html')
    return HTTPException(status_code=404, detail="login.html not found")

@app.get("/signup")
async def serve_signup():
    if os.path.exists("signup.html"): return FileResponse('signup.html')
    return HTTPException(status_code=404, detail="signup.html not found")

@app.get("/profile")
async def serve_profile():
    if os.path.exists("profile.html"): return FileResponse('profile.html')
    return HTTPException(status_code=404, detail="profile.html not found")

@app.get("/calculator")
async def serve_hydraulic(): return FileResponse('calculator.html')

@app.get("/qmax")
async def serve_qmax(): return FileResponse('qmax_calculator.html')
    
@app.get("/load_distribution")
async def serve_load(): return FileResponse('load_distribution_calculator.html')

@app.get("/tractive_effort")
async def serve_te(): return FileResponse('tractive_effort_calculator.html')

@app.get("/vehicle_performance")
async def serve_vp(): return FileResponse('vehicle_performance_calculator.html')

@app.get("/braking")
async def serve_braking():
    if os.path.exists("braking.html"): return FileResponse('braking.html')
    raise HTTPException(status_code=404, detail="braking.html not found")

@app.get("/Diagram.png")
async def serve_diagram_image():
    if os.path.exists("Diagram.png"): return FileResponse('Diagram.png')
    raise HTTPException(status_code=404, detail="Diagram.png not found")

@app.get("/logo.png")
async def serve_logo():
    if os.path.exists("logo.png"): return FileResponse('logo.png')
    raise HTTPException(status_code=404, detail="Logo not found")


# ==========================================
#      CALCULATION ENDPOINTS
# ==========================================

# --- Hydraulic ---
@app.post("/calculate")
async def handle_calculation(raw_input: HydraulicRawInput):
    try:
        inputs, inputs_raw = process_and_validate_hydraulic_inputs(raw_input)
        hydraulic_calculator.inputs_raw = inputs_raw
        if inputs['calc_mode'] == "calc_cc":
            results = hydraulic_calculator.perform_displacement_calculation(inputs)
            report = hydraulic_calculator._generate_mode1_report(inputs, results)
        else:
            results = hydraulic_calculator.perform_speed_calculation(inputs)
            report = hydraulic_calculator._generate_mode2_report(inputs, results)
        return {"report": report, "results": results}
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@app.post("/download_report")
async def download_report(raw_input: HydraulicRawInput):
    try:
        inputs, inputs_raw = process_and_validate_hydraulic_inputs(raw_input)
        hydraulic_calculator.inputs_raw = inputs_raw
        doc = docx.Document()
        filename = "Report.docx"
        if inputs['calc_mode'] == "calc_cc":
            results = hydraulic_calculator.perform_displacement_calculation(inputs)
            hydraulic_calculator._create_mode1_docx(doc, inputs, results)
        else:
            results = hydraulic_calculator.perform_speed_calculation(inputs)
            hydraulic_calculator._create_mode2_docx(doc, inputs, results)
        file_stream = io.BytesIO(); doc.save(file_stream); file_stream.seek(0)
        return StreamingResponse(file_stream, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f"attachment; filename={filename}"})
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

# --- Qmax ---
@app.post("/calculate_qmax")
async def handle_qmax_calculation(raw_input: QmaxRawInput):
    try:
        inputs, inputs_raw = process_and_validate_qmax_inputs(raw_input)
        results = qmax_calculator.perform_calculations(d=inputs['d'], sigma_b=inputs['sigma_b'], v_head=inputs['v_head'])
        report = qmax_calculator.format_detailed_steps(results)
        return {"report": report, "results": results}
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@app.post("/download_qmax_report")
async def download_qmax_report(raw_input: QmaxRawInput):
    try:
        inputs, inputs_raw = process_and_validate_qmax_inputs(raw_input)
        results = qmax_calculator.perform_calculations(d=inputs['d'], sigma_b=inputs['sigma_b'], v_head=inputs['v_head'])
        file_stream = qmax_calculator.create_report_docx(results)
        return StreamingResponse(file_stream, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": "attachment; filename=Qmax_Report.docx"})
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

# --- Load Distribution ---
@app.post("/calculate_load_distribution")
async def handle_load_distribution(raw_input: LoadDistroRawInput):
    try:
        inputs, _ = process_and_validate_load_distro_inputs(raw_input)
        results = perform_load_distro_calc(**inputs)
        report = format_load_distro_steps(inputs, results)
        return {"report": report, "results": results}
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@app.post("/download_load_distribution_report")
async def download_load_distribution_report(raw_input: LoadDistroRawInput):
    try:
        inputs, _ = process_and_validate_load_distro_inputs(raw_input)
        results = perform_load_distro_calc(**inputs)
        file_stream = create_load_distro_docx(inputs, results)
        return StreamingResponse(file_stream, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": "attachment; filename=Load_Dist_Report.docx"})
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

# --- Tractive Effort ---
@app.post("/calculate_tractive_effort")
async def handle_tractive_effort(raw_input: TractiveEffortRawInput):
    try:
        inputs, _ = process_and_validate_te_inputs(raw_input)
        results = perform_te_calculations(inputs)
        report = format_te_report_text(inputs, results)
        return {"report": report, "results": results}
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@app.post("/download_tractive_effort_report")
async def download_te_report(raw_input: TractiveEffortRawInput):
    try:
        inputs, _ = process_and_validate_te_inputs(raw_input)
        results = perform_te_calculations(inputs)
        file_stream = create_te_report_docx(inputs, results)
        return StreamingResponse(file_stream, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": "attachment; filename=Tractive_Report.docx"})
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

# --- Vehicle Performance ---
@app.post("/calculate_performance")
async def handle_performance(raw_input: VehiclePerformanceRawInput):
    try:
        inputs, _ = process_and_validate_vehicle_performance_inputs(raw_input)
        calculator = VehiclePerformanceCalculator(inputs)
        results = calculator.run_tractive_calculation()
        plot_data = calculator.calculate_plot_data()
        table = calculator.calculate_speed_for_shunting_load()
        detailed_calcs = calculator.calculate_detailed_step_by_step()
        return {
            "traction_snapshot": results, 
            "tractive_effort_graph": plot_data["tractive_effort_plot"], 
            "shunting_capability_graph": plot_data["shunting_capability_plot"], 
            "speed_vs_slope_table": table,
            "detailed_calculations": detailed_calcs
        }
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@app.post("/calculate_performance_detailed")
async def handle_performance_detailed(raw_input: VehiclePerformanceRawInput):
    """Return only the detailed step-by-step calculations for documentation purposes."""
    try:
        inputs, _ = process_and_validate_vehicle_performance_inputs(raw_input)
        calculator = VehiclePerformanceCalculator(inputs)
        detailed_calcs = calculator.calculate_detailed_step_by_step()
        return {"detailed_calculations": detailed_calcs}
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@app.post("/download_performance_report")
async def download_performance_report(raw_input: VehiclePerformanceRawInput):
    try:
        inputs, _ = process_and_validate_vehicle_performance_inputs(raw_input)
        calculator = VehiclePerformanceCalculator(inputs)
        file_stream = calculator.create_report_docx()
        return StreamingResponse(file_stream, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": "attachment; filename=Performance_Report.docx"})
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

# --- BRAKING ENDPOINTS ---
@app.post("/braking/calculate")
async def calculate_braking(raw_input: BrakingWebInput):
    """
    Calculate braking performance with full desktop logic replication.
    Returns results table and summary data.
    """
    try:
        inputs = raw_input.dict()
        results = braking_calculator.convert(inputs)
        
        return {
            'results_table': results['results_table'],
            'max_braking_force': results['max_braking_force'],
            'gbr': results['gbr']
        }
    except Exception as e:
        print(f"Braking Calculation Error: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/braking/pdf")
async def download_braking_pdf(raw_input: BrakingWebInput):
    """
    Generate PDF report using LaTeX template.
    Returns ZIP file containing: PDF, TEX source, and all images used.
    """
    try:
        inputs = raw_input.dict()
        
        # Run calculations to populate pdf_placeholders
        results = braking_calculator.convert(inputs)
        placeholders = results['pdf_placeholders']
        
        # Render LaTeX template
        base_dir = os.path.dirname(os.path.abspath(__file__))
        env = Environment(loader=FileSystemLoader(base_dir))
        template = env.get_template('template.tex')
        rendered = template.render(placeholders)
        
        # Create temporary directory for LaTeX compilation
        import tempfile, shutil, zipfile
        from datetime import datetime
        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        workdir = tempfile.mkdtemp(prefix=f"braking_{ts}_")
        
        # Save TEX file
        tex_file = os.path.join(workdir, 'Braking_Report.tex')
        with open(tex_file, 'w', encoding='utf-8') as f:
            f.write(rendered)
        
        # Copy logo and any other images to workdir
        logo_src = os.path.join(base_dir, 'logo.png')
        if os.path.exists(logo_src):
            shutil.copy(logo_src, os.path.join(workdir, 'logo.png'))
        
        # Check for gradient.jpg, curve.jpg, etc.
        image_files = ['gradient.jpg', 'curve.jpg', 'superelevation.png', 'cant.png', 'gauge.jpg']
        for img_name in image_files:
            img_src = os.path.join(base_dir, img_name)
            if os.path.exists(img_src):
                shutil.copy(img_src, os.path.join(workdir, img_name))
        
        pdf_generated = False
        try:
            # Run pdflatex twice to resolve references
            proc1 = subprocess.run(
                ['pdflatex', '-interaction=nonstopmode', 'Braking_Report.tex'],
                cwd=workdir,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                timeout=60
            )
            proc2 = subprocess.run(
                ['pdflatex', '-interaction=nonstopmode', 'Braking_Report.tex'],
                cwd=workdir,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                timeout=30
            )
            
            pdf_path = os.path.join(workdir, 'Braking_Report.pdf')
            if os.path.exists(pdf_path):
                pdf_generated = True
        except Exception as tex_error:
            print(f"LaTeX Error: {tex_error}")
        
        # Create ZIP file with all contents
        zip_path = os.path.join(workdir, f'Braking_Report_{ts}.zip')
        with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
            # Add TEX file
            zipf.write(tex_file, 'Braking_Report.tex')
            
            # Add PDF if generated
            if pdf_generated:
                zipf.write(os.path.join(workdir, 'Braking_Report.pdf'), 'Braking_Report.pdf')
            
            # Add all images that exist
            if os.path.exists(os.path.join(workdir, 'logo.png')):
                zipf.write(os.path.join(workdir, 'logo.png'), 'logo.png')
            
            for img_name in image_files:
                img_path = os.path.join(workdir, img_name)
                if os.path.exists(img_path):
                    zipf.write(img_path, img_name)
            
            # Add auxiliary LaTeX files if they exist
            for ext in ['.aux', '.log']:
                aux_file = os.path.join(workdir, f'Braking_Report{ext}')
                if os.path.exists(aux_file):
                    zipf.write(aux_file, f'Braking_Report{ext}')
        
        # Read ZIP file and return
        with open(zip_path, 'rb') as f:
            zip_bytes = f.read()
        
        shutil.rmtree(workdir, ignore_errors=True)
        
        return StreamingResponse(
            io.BytesIO(zip_bytes),
            media_type="application/zip",
            headers={"Content-Disposition": f"attachment; filename=Braking_Report_{ts}.zip"}
        )
            
    except Exception as e:
        print(f"Braking PDF Error: {e}")
        raise HTTPException(status_code=500, detail=str(e))


if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", 8000))
    print(f"FastAPI server is running on port {port}")
    uvicorn.run("main:app", host="0.0.0.0", port=port, reload=False)